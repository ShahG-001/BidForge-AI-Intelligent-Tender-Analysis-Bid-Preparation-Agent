"""Extract readable text from tender and company evidence files."""
from io import BytesIO


SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt", "md", "csv"}


def read_uploaded_file(uploaded_file) -> tuple[str, str]:
    """Return source-labelled extracted text plus an optional extraction note."""
    filename = uploaded_file.name
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension not in SUPPORTED_EXTENSIONS:
        return "", f"Unsupported file type for {filename}. Use PDF, DOCX, TXT, MD, or CSV."
    try:
        data = uploaded_file.getvalue()
        if extension in {"txt", "md"}:
            return f"[SOURCE: {filename}]\n{data.decode('utf-8', errors='replace')}", ""
        if extension == "csv":
            lines = data.decode("utf-8", errors="replace").splitlines()
            return "\n".join(f"[SOURCE: {filename} | ROW {index}] {line}" for index, line in enumerate(lines, start=1)), ""
        if extension == "pdf":
            from pypdf import PdfReader

            pages = PdfReader(BytesIO(data)).pages
            extracted = [f"[SOURCE: {filename} | PAGE {index}]\n{page.extract_text() or ''}" for index, page in enumerate(pages, start=1)]
            text = "\n\n".join(extracted)
            if not any((page.extract_text() or "").strip() for page in pages):
                return "", f"{filename} appears scanned and has no selectable text. OCR is not included in this starter version; upload a searchable PDF or paste the relevant text."
            return text, ""
        from docx import Document

        document = Document(BytesIO(data))
        blocks = [f"[SOURCE: {filename} | PARAGRAPH {index}] {paragraph.text.strip()}" for index, paragraph in enumerate(document.paragraphs, start=1) if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            blocks.append(f"[SOURCE: {filename} | TABLE {table_index}]")
            blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(blocks), ""
    except Exception as error:
        return "", f"Could not read {filename}: {error}"
